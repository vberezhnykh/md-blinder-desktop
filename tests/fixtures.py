"""Синтетические DOCX, XLSX и PDF для тестов. Никаких реальных данных."""

from __future__ import annotations

import zipfile
from pathlib import Path
from xml.sax.saxutils import escape

import pymupdf

LONG_TEXT = (
    "Synthetic document text with enough body characters to prove that its text layer "
    "contains useful searchable content."
)


def make_pdf(path: Path, pages: list[dict]) -> Path:
    """pages: [{"lines": [(x, y_top_down, text, size)], "image": bool}]"""
    doc = pymupdf.open()
    for spec in pages:
        page = doc.new_page(width=600, height=800)
        if spec.get("image"):
            pix = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 60, 80), 0)
            pix.set_rect(pix.irect, (200, 200, 200))
            page.insert_image(page.rect, pixmap=pix)
        for x, y, text, size in spec.get("lines", []):
            page.insert_text((x, y), text, fontsize=size, fontname="helv")
    doc.save(path)
    doc.close()
    return path


def draw_table(page, x0, y0, widths, texts, row_height=24, merged_below=()):
    """Таблица с рамками. merged_below — ячейки (строка, столбец) без нижней линии: объединение по вертикали."""
    xs = [x0]
    for width in widths:
        xs.append(xs[-1] + width)
    for r, row in enumerate(texts):
        top, bottom = y0 + r * row_height, y0 + (r + 1) * row_height
        for c, text in enumerate(row):
            if r == 0:
                page.draw_line((xs[c], top), (xs[c + 1], top), width=0.7)
            if (r, c) not in merged_below:
                page.draw_line((xs[c], bottom), (xs[c + 1], bottom), width=0.7)
            page.draw_line((xs[c], top), (xs[c], bottom), width=0.7)
            page.draw_line((xs[c + 1], top), (xs[c + 1], bottom), width=0.7)
            if text:
                page.insert_text((xs[c] + 4, top + 16), text, fontsize=10, fontname="helv")


def make_docx(path: Path, body_xml: str) -> Path:
    content_types = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
        '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
        '<Default Extension="xml" ContentType="application/xml"/>'
        '<Override PartName="/word/document.xml" '
        'ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>'
        "</Types>"
    )
    rels = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" '
        'Target="word/document.xml"/></Relationships>'
    )
    document = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        f"<w:body>{body_xml}</w:body></w:document>"
    )
    with zipfile.ZipFile(path, "w") as zf:
        zf.writestr("[Content_Types].xml", content_types)
        zf.writestr("_rels/.rels", rels)
        zf.writestr("word/document.xml", document)
    return path


def para(text: str) -> str:
    return f'<w:p><w:r><w:t xml:space="preserve">{escape(text)}</w:t></w:r></w:p>'


def table(rows: list[list[str]]) -> str:
    cells = "".join(
        "<w:tr>" + "".join(f"<w:tc>{para(c)}</w:tc>" for c in row) + "</w:tr>" for row in rows
    )
    return f"<w:tbl>{cells}</w:tbl>"


def make_xlsx(path: Path, sheets: dict[str, dict[str, object]], hidden: tuple[str, ...] = ()) -> Path:
    """sheets: {имя листа: {"A1": значение}}; формат ячейки задается парой (значение, number_format)."""
    import openpyxl

    book = openpyxl.Workbook()
    book.remove(book.active)
    for name, cells in sheets.items():
        ws = book.create_sheet(name)
        for ref, value in cells.items():
            if isinstance(value, tuple):
                ws[ref] = value[0]
                ws[ref].number_format = value[1]
            else:
                ws[ref] = value
        if name in hidden:
            ws.sheet_state = "hidden"
    book.save(path)
    return path
